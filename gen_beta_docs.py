# -*- coding: utf-8 -*-
"""生成两份交付文档：
1. 报警回传测试步骤（两页）
2. 检测服务使用说明（流程 + 可调整参数）
后续修订后重跑本脚本即可更新。
"""
import subprocess

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT1 = '/home/nvidia/2026/detector/报警回传测试步骤_v1.0.docx'
OUT2 = '/home/nvidia/2026/detector/检测服务使用说明_v1.0.docx'


def device_ip():
    """自动探测本机 IP：在哪台设备上运行，文档里就写哪台的地址"""
    try:
        return subprocess.check_output(["hostname", "-I"]).decode().split()[0]
    except Exception:
        return "127.0.0.1"


IP = device_ip()


def new_doc():
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.line_spacing = 1.25
    for name, size in [('Heading 1', 14), ('Heading 2', 12)]:
        st = doc.styles[name]
        st.font.name = 'Arial'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    return doc


def para(doc, text, bold=False, size=10.5, align=None, space_after=5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = 'Times New Roman'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def mono(doc, text, size=9):
    """命令/文件名等宽显示"""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    return p


def make_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.name = 'Times New Roman'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        shd = cell._tc.get_or_add_tcPr().makeelement(
            qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'D9E2F3'})
        cell._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9)
            r.font.name = 'Times New Roman'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


# ============================================================
# 文档 1：报警回传测试步骤（两页）
# ============================================================
doc = new_doc()
para(doc, '报警回传测试步骤', bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(doc, '罐笼图像检测服务 v1.0（beta1）  2026-10-16', size=9,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

doc.add_heading('一、测试准备', level=1)
para(doc, '1. 服务自检（vsftpd 与 cage-detector 应返回 active；llama-server 精筛关闭时 inactive 属正常）：')
mono(doc, 'systemctl is-active vsftpd cage-detector llama-server')
para(doc, '2. 多模态质检开关（默认关闭；热生效、无需重启服务，llama-server 随开关按需启停）：')
mono(doc, 'curl -X POST http://127.0.0.1:8082/params -d \'{"qc":"off"}\'   # 关闭大模型复检（默认）\n'
          'curl -X POST http://127.0.0.1:8082/params -d \'{"qc":"on"}\'    # 按需开启（自动拉起 llama-server，\n'
          '                                                     # 首次加载约 1~2 分钟，期间报警正常推送）')
para(doc, '3. 启动本机模拟接收端（仅自测用，监听 8081，收到的文件存 /tmp/recv/）：')
mono(doc, 'cd /home/nvidia/2026/detector && python3 recv_test.py &')

doc.add_heading('二、执行测试', level=1)
para(doc, '步骤 1：将测试图片按协议命名后上传至 FTP。文件名规则（六字段下划线分隔）：'
          '17位时间_3位RobotID_3位TaskID_4位相机角度_4位编码器位置_10位预留.jpg，'
          '示例 20261016103000123_007_042_0900_1234_0000000000.jpg。')
mono(doc, f'curl -T 测试图.jpg ftp://{IP}/images/ --user guanlong_ftp:<FTP密码>')
para(doc, '步骤 2：等待处理。质检关闭或无缺陷约 1 秒/张；质检开启时约 6~8 秒/张'
          '（首次质检需等模型加载约 1~2 分钟）。')
para(doc, '步骤 3：按下方对照表核对五项结果（回传报文格式核对：接收端自动存档最近一次'
          '推送的完整 HTTP 请求原文到 /tmp/recv/last_request.txt，cat 即可对照协议 v1.0 逐项核对）。')

doc.add_heading('三、预期结果对照表', level=1)
make_table(doc,
    ['检查项', '位置', '预期结果'],
    [
        ['带框结果图', '/data/ftp/results/', '与原图同名，含检测框与类别标注'],
        ['报警记录', '/data/ftp/alarms/', '同名 .json，含事件编码、置信度、大模型复检结论（场景描述+异常类别）'],
        ['回传推送文件', '接收端 /tmp/recv/', '七字段命名：原名_G002.jpg（事件编码依次追加）'],
        ['回传报文原文', '/tmp/recv/last_request.txt', 'POST /api/ai/alarmevent、multipart name="alarmevent"、Content-Type: application/jpg'],
        ['原图', '/data/ftp/images/', '处理完成后自动删除'],
    ], widths=[3.2, 5.2, 7.6])

doc.add_heading('四、常见问题排查', level=1)
make_table(doc,
    ['现象', '原因', '处理方法'],
    [
        ['推送失败', '接收端未启动，或回传地址未切换',
         '确认 8081 监听；正式环境在 config.py 的 HTTP_SEND["url"] 填实际地址后重启服务'],
        ['报警无质检结论', '质检开关处于 off', 'curl -X POST http://127.0.0.1:8082/params -d \'{"qc":"on"}\''],
        ['llama-server 显示 inactive', '精筛关闭时按需停用，属正常', '无需处理；开精筛后会自动拉起'],
        ['images 目录文件不处理', '扩展名非 .jpg，或上传未完成（.uploading）',
         '检查文件扩展名；等待上传完成'],
        ['读图失败', '文件损坏', '自动移入 /data/ftp/failed/ 留查'],
    ], widths=[3.6, 5.0, 7.4])
para(doc, '服务日志：/home/nvidia/2026/detector/logs/pipeline.log', size=9)

doc.save(OUT1)
print('已生成', OUT1)

# ============================================================
# 文档 2：检测服务使用说明（流程 + 可调整参数）
# ============================================================
doc = new_doc()
para(doc, '罐笼图像检测服务使用说明', bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(doc, '服务组成 / 使用流程 / 可调整参数   2026-10-16', size=9,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

doc.add_heading('一、服务组成', level=1)
make_table(doc,
    ['服务', '作用', '端口'],
    [
        ['vsftpd', '接收机器人 FTP 上传的巡检图片', '21'],
        ['llama-server', 'Qwen3-VL-2B 多模态质检（llama.cpp，LLM-GPU/ViT-CPU 混合推理）', '8080'],
        ['cage-detector', '检测主服务：YOLO26 TRT fp16 推理、结果落盘、报警推送', '—'],
    ], widths=[3.4, 10.2, 1.8])
para(doc, 'vsftpd 与 cage-detector 已配置开机自启与崩溃自动拉起；llama-server 为按需启停'
          '（精筛关闭时自动停止、释放约 1.3GB 显存/内存，开启时自动拉起，首次模型加载约 1~2 分钟，'
          '期间报警按 fail-open 正常推送不漏报）。', size=9.5)

doc.add_heading('二、使用流程', level=1)
para(doc, '机器人按协议命名上传 .jpg 至 /data/ftp/images/ → 服务自动按时间戳最旧优先逐张处理：'
          'TensorRT 推理 → 结果图存 results/ → 检出缺陷时（可选）大模型复检 → '
          '按协议 v1.0 以 HTTP multipart 推送带标记结果图至上位机（原名+事件编码命名）→ '
          '报警 JSON 落盘 alarms/ → 删除原图。无缺陷图片不推送、仅存结果图。')
para(doc, '多模态复检为独立开关（默认关闭、按需开启，record 模式：复检结论记入报警 JSON 供人工复核，'
          '不影响推送）。复检采用固定问答模板：提问"这个图片中存在什么矿井罐笼设施异常"，'
          '回答格式为"图片场景描述：……，图片存在的异常：……"，异常类别限定协议 15 种缺陷类型'
          '与未发现异常共 16 种。运行中通过统一控制接口随时调整，无需重启：')
mono(doc, 'curl http://127.0.0.1:8082/status                                      # 查看运行参数\n'
          'curl -X POST http://127.0.0.1:8082/params -d \'{"qc":"off"}\'           # 关闭复检（默认）\n'
          'curl -X POST http://127.0.0.1:8082/params -d \'{"conf_th":0.4}\'        # 热调置信度阈值\n'
          'curl -X POST http://127.0.0.1:8082/params -d \'{"qc_temperature":0.1}\' # 热调复检生成温度')

doc.add_heading('三、目录结构', level=1)
make_table(doc,
    ['目录', '内容'],
    [
        ['/data/ftp/images/', '巡检原图入口（FTP 落地，处理后自动删除）'],
        ['/data/ftp/results/', '带框结果图（与原图同名）'],
        ['/data/ftp/alarms/', '报警记录 JSON（永久保留，回传失败可据此补发）'],
        ['/data/ftp/failed/', '无法解析的损坏文件'],
        ['~/2026/detector/logs/', '服务运行日志（pipeline.log）'],
    ], widths=[5.5, 9.9])

doc.add_heading('四、统一控制接口与常用运维命令', level=1)
para(doc, '检测服务内置本地控制接口（127.0.0.1:8082），运行参数的查询与热调整统一走此接口，无需重启服务：')
make_table(doc,
    ['操作', '命令'],
    [
        ['查看服务状态', 'systemctl status cage-detector llama-server'],
        ['查看运行日志', 'tail -f ~/2026/detector/logs/pipeline.log'],
        ['查看运行参数', 'curl http://127.0.0.1:8082/status'],
        ['质检开关（热生效）', 'curl -X POST http://127.0.0.1:8082/params -d \'{"qc":"on|off"}\'（联动 llama-server 启停）'],
        ['热调推理阈值', 'curl -X POST http://127.0.0.1:8082/params -d \'{"conf_th":0.4,"iou_th":0.5}\''],
        ['热调质检参数', 'curl -X POST http://127.0.0.1:8082/params -d \'{"qc_timeout":60,"qc_resize":640,"qc_temperature":0.1}\''],
        ['重启检测服务', 'sudo systemctl restart cage-detector'],
    ], widths=[4.5, 10.9])

doc.add_heading('五、可调整参数', level=1)
para(doc, '参数分两级：运行时热调（走控制接口，即时生效，重启后回退 config.py 默认值，质检开关除外——'
          '其状态持久化在开关文件中）与配置文件级（改 config.py 后需重启服务）。'
          '所有参数的可见源头均为 config.py，均带中文注释。', size=9.5)
make_table(doc,
    ['参数', '默认值', '说明', '调整方式'],
    [
        ['qc（质检开关）', 'off', '是否介入大模型复检；三层控制：开关文件/控制接口即时，config.QC_ENABLE 为默认策略', '接口热调，持久化'],
        ['conf_th', '0.35', '报警置信度阈值，调高降误报、调低降漏报', '接口热调 / 改配置重启'],
        ['iou_th', '0.45', '重叠框合并阈值', '接口热调 / 改配置重启'],
        ['qc_timeout', '60s', '单次复检超时', '接口热调 / 改配置重启'],
        ['qc_max_tokens / qc_resize', '48 / 640', '复检输出长度 / 送检图片分辨率', '接口热调 / 改配置重启'],
        ['qc_temperature', '0.1', '复检生成温度（0~1，越低越稳定）', '接口热调 / 改配置重启'],
        ['QC_ON_DEMAND', 'True', 'llama-server 按需启停（精筛关=停服务省资源，开=自动拉起）', '改配置重启'],
        ['SEND_MODE', '"http"', '回传方式：http（协议 v1.0）/ ftp（旧方案回退）', '改配置重启'],
        ['HTTP_SEND["url"]', '本机测试端', '上位机实际接收地址，唯一待切换项', '改配置重启'],
        ['HTTP_SEND["timeout"/"max_retry"]', '15s / 1次', '推送超时与重试', '改配置重启'],
        ['QC_MODE', '"record"', 'record=仅记录不影响推送；gate=复检否决则不推送', '改配置重启'],
        ['EVENT_CODES', '15类映射', '模型类别→G码；新算法类别接入只需增加映射行', '改配置重启'],
        ['POLL_INTERVAL', '1.0s', '目录空闲时轮询间隔', '改配置重启'],
    ], widths=[4.2, 2.4, 6.4, 3.4])

doc.add_heading('六、注意事项', level=1)
for t in [
    '检测链路为单线程串行、最旧优先，请勿改为并发（统一内存平台并发易触发 OOM）。',
    '回传失败不阻断主链路：报警 JSON 永久保留在 alarms/，可据此补发。',
    '修改 config.py 后需 systemctl restart cage-detector 生效；质检开关除外。',
    '精筛开启后首次质检有约 1~2 分钟模型加载等待，期间报警按 fail-open 正常推送不漏报。',
    'llama-server 按需启停依赖 /etc/sudoers.d/llama-ondemand（放行 systemctl start/stop llama-server 免密）。',
    '算法迭代新增类别（如卡子类 G006/G008）：新权重导出 engine 后，在 CLASS_NAMES 与 EVENT_CODES 各加一行即可接入。',
]:
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(t)
    r.font.size = Pt(9.5)
    r.font.name = 'Times New Roman'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

doc.save(OUT2)
print('已生成', OUT2)
